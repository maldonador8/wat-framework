#!/usr/bin/env python3
"""Converts the optimized resume markdown to a Word .docx file with proper ATS-friendly styles."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def setup_heading_styles(doc):
    """Configure Heading 1, 2, 3 styles for ATS recognition."""
    # Heading 1 — Name
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Heading 2 — Section headers (PROFESSIONAL EXPERIENCE, SKILLS, etc.)
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 51, 102)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)

    # Heading 3 — Job titles
    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)


def add_bullet_with_bold_label(doc, text):
    """Add a bullet point, bolding the category label before the colon if present."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)

    # Check if line has a bold label pattern like "Cloud - AWS: stuff"
    match = re.match(r"^(.+?:)\s*(.+)$", text)
    if match and len(match.group(1)) < 30:
        # Bold the label
        run_label = p.add_run(match.group(1) + " ")
        run_label.bold = True
        run_label.font.size = Pt(10)
        run_label.font.name = "Calibri"
        # Normal text for the rest
        run_rest = p.add_run(match.group(2))
        run_rest.font.size = Pt(10)
        run_rest.font.name = "Calibri"
    else:
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.name = "Calibri"

    return p


def build_docx(md_path, output_path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Set up proper heading styles for ATS
    setup_heading_styles(doc)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    lines = content.split("\n")
    i = 0
    in_skills_or_certs = False

    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines, horizontal rules, and the note at the bottom
        if not line or (line.startswith("---") and "|" not in line) or line.startswith("*NOTE:") or line.startswith("*[X]"):
            i += 1
            continue

        # H1 — Name (use proper Heading 1 style)
        if line.startswith("# ") and not line.startswith("## "):
            doc.add_paragraph(line[2:].strip(), style="Heading 1")
            i += 1
            continue

        # Contact info — address line
        if re.match(r"^\d+\s+\w+.*\d{5}", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(80, 80, 80)
            run.font.name = "Calibri"
            i += 1
            continue

        # Contact info — phone/email line
        if "|" in line and "@" in line:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(80, 80, 80)
            run.font.name = "Calibri"
            i += 1
            continue

        if re.match(r"^\(?\d{3}\)?[\s-]", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(80, 80, 80)
            run.font.name = "Calibri"
            i += 1
            continue

        # H2 — Section headers (use proper Heading 2 style)
        if line.startswith("## "):
            section_name = line[3:].strip()
            doc.add_paragraph(section_name, style="Heading 2")
            in_skills_or_certs = section_name in ("TECHNICAL SKILLS", "CERTIFICATIONS")
            i += 1
            continue

        # H3 — Job titles (use proper Heading 3 style)
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
            in_skills_or_certs = False
            i += 1
            continue

        # Bold date lines like **06/2025 – Present**
        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = "Calibri"
            run.font.color.rgb = RGBColor(100, 100, 100)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            i += 1
            continue

        # Bullet points
        if line.startswith("- "):
            text = line[2:].strip()
            # Remove markdown formatting
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            text = re.sub(r"\*(.*?)\*", r"\1", text)
            text = re.sub(r"\[(.*?)\]\(.*?\)", r"\1", text)

            # Use bold labels for skills section items
            if in_skills_or_certs and ":" in text and text.index(":") < 25:
                add_bullet_with_bold_label(doc, text)
            else:
                p = doc.add_paragraph(text, style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = "Calibri"
            i += 1
            continue

        # Regular text
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
        text = re.sub(r"\*(.*?)\*", r"\1", text)
        if text.strip():
            p = doc.add_paragraph(text.strip())
            p.paragraph_format.space_after = Pt(2)
        i += 1

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent.parent
    md = project_root / "docs" / "resume_optimized.md"
    out = project_root / "docs" / "Rafael_Maldonado_Resume.docx"
    result = build_docx(md, out)
    print(f"Resume saved to: {result}")
